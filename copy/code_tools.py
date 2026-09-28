"""General-mode tools: run Python in the sandbox and keep the files it makes.

Only General mode binds these (graph.GENERAL_TOOLS, run by their own ToolNode),
so Documents mode can never execute code. The code itself runs in the
code-runner sidecar (docker/code-runner), never in this process.

A tool result has two halves (response_format="content_and_artifact"):
  content   what the model reads back: output and file names, kept small;
  artifact  what the UI shows: the code, its output and the saved file ids.
"""
from __future__ import annotations

import base64
import binascii
import json
import logging
import os
import re
from datetime import timedelta
from typing import Annotated

import httpx
from asgiref.sync import sync_to_async
from django.conf import settings
from django.http import FileResponse, Http404
from django.utils import timezone
from langchain_core.tools import tool
from langgraph.prebuilt import InjectedState
from rest_framework.decorators import api_view

from chat.models import GeneratedFile

logger = logging.getLogger(__name__)

# extension -> (mime, kind)
FILE_TYPES: dict[str, tuple[str, str]] = {
    ".png": ("image/png", "image"),
    ".jpg": ("image/jpeg", "image"),
    ".jpeg": ("image/jpeg", "image"),
    ".svg": ("image/svg+xml", "image"),
    ".pdf": ("application/pdf", "document"),
    ".docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "document"),
    ".xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "spreadsheet"),
    ".pptx": ("application/vnd.openxmlformats-officedocument.presentationml.presentation", "presentation"),
    ".csv": ("text/csv", "data"),
    ".txt": ("text/plain", "data"),
    ".json": ("application/json", "data"),
    ".md": ("text/markdown", "data"),
}
# Shown in the page itself. Everything else, SVG included (it can carry
# script), is served as a download.
INLINE_MIMES = {"image/png", "image/jpeg"}
MODEL_OUTPUT_CHARS = 4000
_UNSAFE_NAME = re.compile(r"[^\w.\- ]+", re.UNICODE)


class RunnerUnavailable(RuntimeError):
    pass


def generated_dir(file_id) -> str:
    return os.path.join(settings.MEDIA_ROOT, "chat", "generated", str(file_id))


def safe_filename(name: str) -> str | None:
    """A plain basename with an allowed extension, or None."""
    base = os.path.basename(name or "").strip()
    stem, extension = os.path.splitext(base)
    extension = extension.lower()
    if extension not in FILE_TYPES:
        return None
    stem = _UNSAFE_NAME.sub("_", stem).strip(" ._")[:120] or "file"
    return stem + extension


def serialize_file(file: GeneratedFile) -> dict:
    return {
        "id": str(file.id),
        "filename": file.filename,
        "kind": file.kind,
        "mime": file.mime,
        "size": file.size,
        "inline": file.mime in INLINE_MIMES,
        "url": f"/api/chat/files/{file.id}/",
    }


async def call_runner(code: str) -> dict:
    """POST the code to the sandbox. Raises RunnerUnavailable when it cannot run."""
    timeout = settings.CODE_RUN_TIMEOUT_S
    try:
        async with httpx.AsyncClient(timeout=timeout + 15) as client:
            response = await client.post(
                f"{settings.CODE_RUNNER_URL.rstrip('/')}/run",
                json={"code": code, "timeout_s": timeout},
            )
    except httpx.HTTPError as exc:
        logger.warning("code runner unreachable: %s", exc)
        raise RunnerUnavailable("the code runner is not reachable") from exc
    if response.status_code == 429:
        raise RunnerUnavailable("the code runner is busy; try again in a moment")
    if response.status_code != 200:
        detail = ""
        try:
            detail = response.json().get("detail", "")
        except ValueError:
            pass
        raise RunnerUnavailable(detail or f"the code runner failed ({response.status_code})")
    return response.json()


@sync_to_async(thread_sensitive=True)
def save_files(files: list[dict], user_id, conversation_id) -> list[GeneratedFile]:
    saved: list[GeneratedFile] = []
    for item in files:
        name = safe_filename(str(item.get("name") or ""))
        if name is None:
            continue
        try:
            data = base64.b64decode(item.get("b64") or "", validate=True)
        except (binascii.Error, ValueError):
            logger.warning("code runner returned undecodable file %r", item.get("name"))
            continue
        mime, kind = FILE_TYPES[os.path.splitext(name)[1]]
        record = GeneratedFile(
            owner_id=user_id, conversation_id=conversation_id,
            kind=kind, filename=name, mime=mime, size=len(data),
        )
        directory = generated_dir(record.id)
        os.makedirs(directory, exist_ok=True)
        record.storage_path = os.path.join(directory, name)
        with open(record.storage_path, "wb") as fh:
            fh.write(data)
        record.save()
        saved.append(record)
    return saved


def _clip(text: str) -> str:
    text = text or ""
    return text if len(text) <= MODEL_OUTPUT_CHARS else text[:MODEL_OUTPUT_CHARS] + "\n[… truncated]"


@tool(response_format="content_and_artifact")
async def run_python(
    code: str,
    user_id: Annotated[int | None, InjectedState("user_id")],
    conversation_id: Annotated[str | None, InjectedState("conversation_id")],
) -> tuple[str, dict]:
    """Run Python 3 code inside the sandbox and return the execution result and any files successfully created.

==================================================
PRIMARY PURPOSE
===============

This tool is an EXECUTION tool.

Use it when the user asks you to:

* perform calculations
* analyze data
* create graphs or charts
* create Excel files
* create Word DOCX files
* create PowerPoint PPTX files
* create PDF files
* generate files using Python
* manipulate or transform data using Python

DO NOT simply return Python code when the user expects the task to be performed.

Write the Python code, execute it, inspect the result, and return the resulting files/output.

If execution fails:

1. Read the traceback.
2. Identify the exact failing line.
3. Identify the type of error.
4. Fix the specific problem.
5. Execute the corrected code again.
6. Continue until the task succeeds or a genuine limitation is reached.

DO NOT blindly execute the same failed code again.

Never tell the user that a file was created unless the Python execution succeeded and the file was actually saved.

==================================================
STRICT API RULE — VERY IMPORTANT
================================

The examples in this tool description are VERIFIED API PATTERNS.

Treat them as an API reference and whitelist.

You may adapt the examples to the user's request.

You MAY change:

* text
* numbers
* data
* colors
* filenames
* dimensions
* positions
* number of rows
* number of columns
* number of slides
* number of paragraphs
* calculations
* chart values
* formatting values

You MAY:

* repeat demonstrated operations
* combine demonstrated operations
* use demonstrated operations in a different order when logically valid
* create larger documents using the same demonstrated patterns
* adapt demonstrated patterns to the user's specific use case

You MUST NOT:

* invent methods
* invent attributes
* guess APIs
* assume APIs exist because they sound reasonable
* copy APIs from another Python library
* mix object models between libraries
* use undocumented shortcuts when a demonstrated pattern exists
* create a new API pattern merely because it seems convenient

IMPORTANT:

The examples define HOW to use the libraries.

They do NOT mean you must copy their exact content.

The user may request a completely different document, presentation, spreadsheet, calculation, or graph.

Adapt the CONTENT while preserving the VERIFIED API PATTERNS.

==================================================
LIBRARY SEPARATION
==================

These libraries have DIFFERENT APIs:

DOCX:
python-docx

PPTX:
python-pptx

Excel:
openpyxl

PDF:
reportlab

Never transfer an API from one library to another.

For example:

DOCX Cell:
cell.paragraphs

PPTX Cell:
cell.text_frame

Excel Cell:
cell.fill

These are different objects and different APIs.

Do not assume that because an attribute exists in one library it exists in another.

==================================================
AVAILABLE LIBRARIES
===================

The sandbox provides these libraries:

* math
* statistics
* numpy
* scipy
* sympy
* pandas
* matplotlib
* openpyxl
* python-docx
* python-pptx
* reportlab
* arabic_reshaper
* bidi

Use only available libraries.

There is no internet access.

Do not use:

* requests
* urllib
* web APIs
* internet downloads
* network services

Do not use input().

Do not use plt.show().

Save graphs using plt.savefig().

==================================================
OBJECT HIERARCHIES
==================

Remember these object hierarchies exactly.

DOCX:

Document
→ Paragraph
→ Run

DOCX table:

Table
→ Cell
→ Paragraph
→ Run

PPTX:

Presentation
→ Slide
→ Shape
→ TextFrame
→ Paragraph
→ Run

PPTX table:

Table
→ Cell
→ TextFrame
→ Paragraph
→ Run

Excel:

Workbook
→ Worksheet
→ Cell

Do not skip hierarchy levels by guessing.

==================================================
EMPTY COLLECTION RULE
=====================

Never assume a collection contains an element.

For example:

```
paragraph.runs[0]
```

can fail if the paragraph has no runs.

When creating formatted text, prefer explicitly creating the run:

```
paragraph = doc.add_paragraph()
run = paragraph.add_run("Text")
```

Then format the run.

Do NOT write:

```
paragraph.paragraphs[0]
```

because a Paragraph does not have a paragraphs attribute.

Do NOT write:

```
paragraph.paragraphs[0].add_run()
```

==================================================
DOCX — PYTHON-DOCX
==================

Import:

from docx import Document
from docx.shared import RGBColor, Pt

Create a document:

doc = Document()

Add a heading:

doc.add_heading("Demo Document", level=1)

Create a paragraph:

paragraph = doc.add_paragraph()

Create and format a run:

run = paragraph.add_run("Hello World")

run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = RGBColor(11, 77, 66)

Save:

doc.save("example.docx")

---

## DOCX — MULTIPLE PARAGRAPHS

Example:

doc = Document()

p1 = doc.add_paragraph("First paragraph.")

p2 = doc.add_paragraph()
run = p2.add_run("Second paragraph.")
run.bold = True

p3 = doc.add_paragraph("Third paragraph.")

doc.save("example.docx")

IMPORTANT:

`p2` is already a Paragraph.

WRONG:

p2.paragraphs[0].add_run()

CORRECT:

p2.add_run("Text")

---

## DOCX — RUNS

A Paragraph contains Runs.

Correct:

paragraph = doc.add_paragraph()
run = paragraph.add_run("Hello")
run.bold = True

If text already exists:

paragraph = doc.add_paragraph("Hello")
run = paragraph.runs[0]
run.bold = True

Only use `.runs[0]` when a run is guaranteed to exist.

Do not access `.runs[0]` on an empty newly-created paragraph.

---

## DOCX — TABLE

Create:

table = doc.add_table(rows=3, cols=2)

Set cell text:

table.cell(0, 0).text = "Name"
table.cell(0, 1).text = "Value"

table.cell(1, 0).text = "Apple"
table.cell(1, 1).text = "100"

table.cell(2, 0).text = "Orange"
table.cell(2, 1).text = "200"

Save:

doc.save("example_table.docx")

---

## DOCX — TABLE CELL FORMATTING

Correct hierarchy:

Cell
→ Paragraph
→ Run

Example:

cell = table.cell(0, 0)

paragraph = cell.paragraphs[0]

run = paragraph.add_run("Example")

run.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(11, 77, 66)

IMPORTANT:

A DOCX `_Cell` does NOT have:

```
cell.runs
```

A DOCX `_Cell` does NOT have:

```
cell.fill
```

WRONG:

cell.runs[0]

WRONG:

cell.fill = RGBColor(11, 77, 66)

Use:

cell.paragraphs[0]

then:

paragraph.add_run("Text")

If text has already been assigned:

cell.text = "Example"

paragraph = cell.paragraphs[0]

run = paragraph.runs[0]

This is valid because assigning `cell.text` creates the required text structure.

---

## DOCX — TABLE ITERATION

To iterate:

for row in table.rows:
for cell in row.cells:
print(cell.text)

Do NOT use:

table.rows.count_rows()

Do NOT use:

table.columns.count_columns()

Do NOT use:

table.columns["A"]

Do NOT use:

table.rows["1"]

DOCX tables do not use Excel-style column names.

==================================================
PPTX — PYTHON-PPTX
==================

Import:

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

Create:

prs = Presentation()

Create a blank slide:

slide = prs.slides.add_slide(prs.slide_layouts[6])

Prefer blank slides with explicit textboxes.

---

## PPTX — TEXTBOX

Create:

textbox = slide.shapes.add_textbox(
Inches(1),
Inches(1),
Inches(8),
Inches(1)
)

Access the text frame:

text_frame = textbox.text_frame

Set text:

text_frame.text = "Hello World"

Access paragraph:

paragraph = text_frame.paragraphs[0]

Access existing run:

run = paragraph.runs[0]

Format:

run.font.size = Pt(24)
run.font.bold = True
run.font.color.rgb = RGBColor(11, 77, 66)

Save:

prs.save("example.pptx")

---

## PPTX — TEXT HIERARCHY

Correct:

slide
→ shape
→ text_frame
→ paragraph
→ run

WRONG:

slide.runs

WRONG:

textbox.runs

WRONG:

text_frame.runs

A Run belongs to a Paragraph.

---

## PPTX — RGB COLORS

RGBColor requires three separate integer values:

RGBColor(11, 77, 66)

WRONG:

RGBColor(0x2C5F7D)

WRONG:

RGBColor("#2C5F7D")

Correct:

RGBColor(red, green, blue)

For example:

RGBColor(44, 95, 125)

---

## PPTX — MULTIPLE TEXTBOXES

Example:

textbox1 = slide.shapes.add_textbox(
Inches(1),
Inches(1),
Inches(8),
Inches(1)
)

textbox1.text_frame.text = "Title"

paragraph = textbox1.text_frame.paragraphs[0]
run = paragraph.runs[0]

run.font.size = Pt(28)
run.font.bold = True

textbox2 = slide.shapes.add_textbox(
Inches(1),
Inches(2.5),
Inches(8),
Inches(2)
)

textbox2.text_frame.text = "Supporting content."

---

## PPTX — TABLE

Create:

table = slide.shapes.add_table(
3,
2,
Inches(1),
Inches(2),
Inches(8),
Inches(2)
).table

Set values:

table.cell(0, 0).text = "Name"
table.cell(0, 1).text = "Value"

table.cell(1, 0).text = "Apple"
table.cell(1, 1).text = "100"

table.cell(2, 0).text = "Orange"
table.cell(2, 1).text = "200"

---

## PPTX — TABLE CELL FORMATTING

Correct hierarchy:

Cell
→ TextFrame
→ Paragraph
→ Run

Example:

cell = table.cell(0, 0)

paragraph = cell.text_frame.paragraphs[0]

run = paragraph.runs[0]

run.font.bold = True
run.font.size = Pt(14)
run.font.color.rgb = RGBColor(11, 77, 66)

Do NOT use:

cell.runs

Do NOT use:

cell.paragraphs

Use:

cell.text_frame.paragraphs

---

## PPTX — TABLE DIMENSIONS

Do NOT invent:

table.rows.count_rows()

table.columns.count_columns()

table.count_rows()

table.count_columns()

If the row and column counts are known because you created the table, use those known values.

---

## PPTX — PLACEHOLDERS

Do NOT assume placeholder indices exist.

WRONG:

slide.placeholders[2]

unless the exact placeholder is known to exist.

Prefer:

slide = prs.slides.add_slide(prs.slide_layouts[6])

Then explicitly create textboxes.

==================================================
EXCEL — OPENPYXL
================

Import:

from openpyxl import Workbook

Create:

wb = Workbook()

ws = wb.active

Write cells:

ws["A1"] = "Name"
ws["B1"] = "Value"

ws["A2"] = "Apple"
ws["B2"] = 100

ws["A3"] = "Orange"
ws["B3"] = 200

Save:

wb.save("example.xlsx")

---

## EXCEL — FORMATTING

Import:

from openpyxl.styles import PatternFill, Font

Example fill:

ws["A1"].fill = PatternFill(
fill_type="solid",
fgColor="0B4D42"
)

Example font:

ws["A1"].font = Font(
color="FFFFFF",
bold=True
)

IMPORTANT:

Excel Cell formatting is NOT DOCX Cell formatting.

WRONG:

cell.fill = RGBColor(11, 77, 66)

Use:

PatternFill(...)

for Excel cell fills.

---

## EXCEL — CALCULATIONS

Example:

ws["A1"] = "Value 1"
ws["B1"] = "Value 2"
ws["C1"] = "Total"

ws["A2"] = 100
ws["B2"] = 200

ws["C2"] = "=A2+B2"

Save:

wb.save("example.xlsx")

---

## EXCEL — ROWS

Valid properties:

ws.max_row
ws.max_column

Example:

print(ws.max_row)
print(ws.max_column)

Iteration:

for row in ws.iter_rows():
for cell in row:
print(cell.value)

Do NOT invent:

ws.count_rows()

ws.count_columns()

==================================================
CALCULATIONS
============

For basic calculations use normal Python.

Example:

a = 10
b = 5

total = a + b
difference = a - b
product = a * b
division = a / b
average = (a + b) / 2

print("Total:", total)
print("Difference:", difference)
print("Product:", product)
print("Division:", division)
print("Average:", average)

---

## MATH

Use:

import math

value = math.sqrt(25)

print(value)

Use demonstrated functions rather than inventing function names.

---

## STATISTICS

Use:

import statistics

values = [10, 20, 30, 40]

mean = statistics.mean(values)
median = statistics.median(values)

print("Mean:", mean)
print("Median:", median)

---

## NUMPY

Use:

import numpy as np

values = np.array([10, 20, 30, 40])

print(values.mean())
print(values.sum())
print(values.max())
print(values.min())

==================================================
SIMPLE CHARTS
=============

Use matplotlib.

Import:

import matplotlib.pyplot as plt

---

## LINE CHART

Example:

x = [1, 2, 3, 4]
y = [10, 20, 15, 30]

plt.plot(x, y)

plt.xlabel("X")
plt.ylabel("Y")
plt.title("Example Graph")

plt.savefig("example_graph.png")

plt.close()

IMPORTANT:

Do NOT use:

plt.show()

Always save the chart.

---

## BAR CHART

Example:

categories = ["A", "B", "C"]
values = [10, 20, 15]

plt.bar(categories, values)

plt.xlabel("Category")
plt.ylabel("Value")
plt.title("Example Bar Chart")

plt.savefig("example_bar.png")

plt.close()

You may change the categories and values.

==================================================
PANDAS
======

Example:

import pandas as pd

data = {
"Name": ["Ali", "Ahmed", "Sara"],
"Score": [90, 85, 95]
}

df = pd.DataFrame(data)

print(df)

print(df["Score"].mean())

You may adapt the data and calculations.

==================================================
PDF — REPORTLAB
===============

Import:

from reportlab.pdfgen import canvas

Create:

pdf = canvas.Canvas("example.pdf")

Set font:

pdf.setFont("Helvetica", 18)

Draw text:

pdf.drawString(
72,
750,
"Hello World"
)

Change font:

pdf.setFont("Helvetica", 12)

pdf.drawString(
72,
720,
"This is a simple PDF."
)

Save:

pdf.save()

---

## PDF — COLOR

Use:

from reportlab.lib import colors

pdf.setFillColor(
colors.HexColor("#0B4D42")
)

pdf.drawString(
72,
750,
"Green Text"
)

Reset:

pdf.setFillColor(colors.black)

---

## PDF — MULTIPLE PAGES

Before starting another page:

pdf.showPage()

Then continue drawing.

Do not invent reportlab methods.

==================================================
FILE RULES
==========

Save files in the current working directory.

Use simple descriptive filenames:

example.docx
example.pptx
example.xlsx
example.pdf
example.png

When creating a requested file:

1. Generate the Python code.
2. Execute it.
3. Inspect stdout and stderr.
4. Fix errors if necessary.
5. Execute again.
6. Confirm successful completion.
7. Ensure the file was saved.
8. Return the resulting file.

Never claim that a file was created based only on the Python source.

==================================================
ERROR RECOVERY
==============

When an error occurs, carefully classify it.

Examples:

IndexError:
An indexed collection may be empty.

AttributeError:
The object may not support that attribute or the wrong hierarchy may be used.

TypeError:
The function may have received the wrong number or type of arguments.

KeyError:
A dictionary/object key may not exist.

IndentationError:
Fix the Python indentation.

NameError:
A variable or import may be missing.

If the error indicates an invalid API, do NOT invent another API.

Replace it with a known demonstrated pattern.

Example:

If this fails:

```
p2.paragraphs[0].add_run()
```

because `p2` is already a Paragraph,

use:

```
p2.add_run("Text")
```

Example:

If this fails:

```
cell.runs[0]
```

for a DOCX table cell,

use:

```
paragraph = cell.paragraphs[0]
run = paragraph.add_run("Text")
```

Example:

If this fails:

```
table.columns["A"]
```

in python-docx,

do not try another Excel-style column operation.

Use the demonstrated python-docx table APIs instead.

==================================================
FINAL PRINCIPLE
===============

Be flexible about WHAT the user wants.

Be strict about HOW the Python library is used.

The user request determines the content and final result.

The verified examples determine the API patterns.

Use:

USER REQUEST
↓
SELECT APPROPRIATE VERIFIED PATTERNS
↓
ADAPT CONTENT
↓
GENERATE CODE
↓
EXECUTE
↓
INSPECT
↓
FIX
↓
EXECUTE AGAIN
↓
VERIFY OUTPUT
↓
RETURN RESULT

Do not sacrifice API correctness for complexity.

A simple working implementation is always preferable to a complicated implementation containing guessed APIs.

    """






    artifact = {"code": code, "ok": False, "stdout": "", "stderr": "", "duration_ms": None, "file_ids": []}
    # Never raise: ToolNode would turn the exception into a message for the
    # model, but the stream would get no on_tool_end, so the UI could not show
    # the run. Every failure comes back as a failed run instead.
    try:
        result = await call_runner(code)
        files = await save_files(result.get("files") or [], user_id, conversation_id) if user_id else []
    except RunnerUnavailable as exc:
        artifact["stderr"] = f"Code execution is unavailable: {exc}."
        return json.dumps({"ok": False, "error": artifact["stderr"]}), artifact
    except Exception:
        logger.exception("run_python failed")
        artifact["stderr"] = "Code execution failed unexpectedly; the result could not be read."
        return json.dumps({"ok": False, "error": artifact["stderr"]}), artifact

    artifact.update(
        ok=bool(result.get("ok")),
        stdout=result.get("stdout") or "",
        stderr=result.get("stderr") or "",
        duration_ms=result.get("duration_ms"),
        file_ids=[str(f.id) for f in files],
    )
    content = {
        "ok": artifact["ok"],
        "stdout": _clip(artifact["stdout"]),
        "stderr": _clip(artifact["stderr"]),
        "files_saved": [f.filename for f in files],
    }
    return json.dumps(content, ensure_ascii=False), artifact


GENERAL_TOOLS = [run_python]


def delete_stale_unbound_files(max_age: timedelta = timedelta(hours=24)) -> int:
    """Files whose turn never finished (the tab closed mid-run) and so never got a message."""
    stale = GeneratedFile.objects.filter(message__isnull=True, created_at__lt=timezone.now() - max_age)
    count = 0
    for record in stale:
        record.delete()  # the post_delete signal removes the file on disk
        count += 1
    return count


# ---- Views -------------------------------------------------------------------

@api_view(["GET"])
def download_generated_file(request, file_id):
    record = GeneratedFile.objects.filter(id=file_id, owner=request.user).first()
    if record is None or not os.path.exists(record.storage_path):
        raise Http404("no such file")
    response = FileResponse(
        open(record.storage_path, "rb"),
        content_type=record.mime,
        as_attachment=record.mime not in INLINE_MIMES,
        filename=record.filename,
    )
    response["X-Content-Type-Options"] = "nosniff"
    return response
