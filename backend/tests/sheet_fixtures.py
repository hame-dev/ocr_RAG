"""A realistic student workbook, built in code.

It has the things real sheets have and clean test data lacks: a title merged
across the top, a blank row, a group header ("Grades") merged over two
columns with the subject names below it, header cells merged down over both
header rows, a class merged down over its students, a section written once
per group with blanks below, Arabic-Indic digits, and Arabic yes/no.
"""
from __future__ import annotations

import io

STUDENTS = [
    # id, first, last, class, math, physics (Arabic-Indic), passed, section
    (20231001, "Ahmed", "Ali", "10-A", 87, "٩٠", "نعم", "A"),
    (20231002, "Sara", "Hassan", None, 45, "٥٠", "لا", None),
    (20231003, "Omar", "Khalid", None, 72, "٦٨", "نعم", None),
    (20231004, "Lina", "Saeed", "10-B", 91, "٩٥", "نعم", "B"),
    (20231005, "Yusuf", "Nasser", None, 38, "٤٠", "لا", None),
    (20231006, "Mona", "Adel", None, 66, "٧٠", "نعم", None),
]

EXPECTED_HEADERS = [
    "Student ID", "First name", "Last name", "Class", "Grades / Math", "Grades / Physics", "Passed", "Section",
]
FIRST_DATA_ROW = 5


def student_workbook() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Grades 2024"

    sheet["A1"] = "School report 2024"
    sheet.merge_cells("A1:H1")
    # Row 2 is blank.
    for column, label in zip("ABCD", ["Student ID", "First name", "Last name", "Class"]):
        sheet[f"{column}3"] = label
        sheet.merge_cells(f"{column}3:{column}4")
    sheet["E3"] = "Grades"
    sheet.merge_cells("E3:F3")
    sheet["E4"], sheet["F4"] = "Math", "Physics"
    for column, label in zip("GH", ["Passed", "Section"]):
        sheet[f"{column}3"] = label
        sheet.merge_cells(f"{column}3:{column}4")

    for offset, student in enumerate(STUDENTS):
        row = FIRST_DATA_ROW + offset
        for column, value in zip("ABCDEFGH", student):
            if value is not None:
                sheet[f"{column}{row}"] = value
    # Each class is written once, merged down over its students.
    sheet.merge_cells("D5:D7")
    sheet.merge_cells("D8:D10")

    # A hidden sheet is never read.
    hidden = workbook.create_sheet("scratch")
    hidden["A1"], hidden["B1"] = "x", "y"
    hidden["A2"], hidden["B2"] = 1, 2
    hidden.sheet_state = "hidden"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def arabic_csv() -> bytes:
    """A CSV as Excel on Arabic Windows saves it: cp1256, not UTF-8."""
    text = "رقم الطالب,اسم الطالب,الدرجة\r\n1,أحمد علي,90\r\n2,سارة حسن,45\r\n3,عمر خالد,72\r\n"
    return text.encode("cp1256")


class FakeProposer:
    """Stands in for Ollama: returns a fixed proposal, or fails."""

    def __init__(self, proposal: dict | None = None, error: Exception | None = None):
        self.proposal = proposal
        self.error = error
        self.prompts: list[str] = []

    def generate_json(self, model, prompt, schema, **kwargs):
        import json

        self.prompts.append(prompt)
        if self.error:
            raise self.error
        return self.proposal, json.dumps(self.proposal)


STUDENT_PROPOSAL = {
    "entity": "student",
    "entity_ar": "طالب",
    "description": "Grades of 6 students in two subjects.",
    "title_template": "{full_name} ({student_id})",
    "columns": [
        {"source": "Student ID", "key": "student_id", "label": "Student ID", "type": "id",
         "role": "identifier", "description": ""},
        {"source": "First name", "key": "first_name", "label": "First name", "type": "text",
         "role": "title", "description": ""},
        {"source": "Last name", "key": "last_name", "label": "Last name", "type": "text",
         "role": "title", "description": ""},
        {"source": "Class", "key": "class", "label": "Class", "type": "category",
         "role": "attribute", "description": ""},
        {"source": "Grades / Math", "key": "math", "label": "Math", "type": "integer",
         "role": "attribute", "description": "math grade"},
        {"source": "Grades / Physics", "key": "physics", "label": "Physics", "type": "integer",
         "role": "attribute", "description": ""},
        {"source": "Passed", "key": "passed", "label": "Passed", "type": "boolean",
         "role": "attribute", "description": ""},
        {"source": "Section", "key": "section", "label": "Section", "type": "text",
         "role": "attribute", "description": ""},
    ],
    "combine": [{"key": "full_name", "label": "Full name", "sources": ["first_name", "last_name"], "sep": " "}],
}
