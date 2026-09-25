from __future__ import annotations

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from langchain_core.messages import HumanMessage
from PIL import Image

from chat.attachments import build_human_content
from chat.models import ChatAttachment, Conversation, Message
from chat.pipeline import text_of
from tests.fakes import Script, stream, use_script

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def _media(settings, tmp_path):
    settings.MEDIA_ROOT = str(tmp_path)


def _png(text_color="black") -> bytes:
    image = Image.new("RGB", (400, 200), "white")
    buf = io.BytesIO()
    image.save(buf, "PNG")
    return buf.getvalue()


def _upload(client, name, data, content_type="application/octet-stream"):
    return client.post(
        "/api/chat/attachments/", data={"file": SimpleUploadedFile(name, data, content_type=content_type)}
    )


def test_image_upload_is_normalized_and_previewable(auth_client, other_client):
    response = _upload(auth_client, "photo.png", _png(), "image/png")

    assert response.status_code == 201, response.content
    body = response.json()
    assert body["kind"] == "image" and body["has_preview"]
    preview = auth_client.get(body["preview_url"])
    assert preview.status_code == 200 and preview["Content-Type"] == "image/jpeg"
    assert other_client.get(body["preview_url"]).status_code == 404


def test_file_type_is_checked_by_content_not_extension(auth_client):
    response = _upload(auth_client, "evil.png", b"<script>not an image</script>", "image/png")

    assert response.status_code == 400
    assert not ChatAttachment.objects.exists()


def test_text_and_docx_are_extracted(auth_client):
    import docx

    document = docx.Document()
    document.add_paragraph("Quarterly revenue grew 12%.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Q1"
    table.rows[0].cells[1].text = "4,250"
    buf = io.BytesIO()
    document.save(buf)

    text = _upload(auth_client, "notes.md", "# Heading\nمرحبا".encode(), "text/markdown").json()
    word = _upload(auth_client, "report.docx", buf.getvalue()).json()

    assert text["kind"] == "text"
    assert ChatAttachment.objects.get(id=text["id"]).extracted_text == "# Heading\nمرحبا"
    extracted = ChatAttachment.objects.get(id=word["id"]).extracted_text
    assert "Quarterly revenue grew 12%." in extracted and "Q1 | 4,250" in extracted


def test_scanned_pdf_becomes_page_images(auth_client):
    buf = io.BytesIO()
    Image.new("RGB", (600, 800), "white").save(buf, "PDF")

    body = _upload(auth_client, "scan.pdf", buf.getvalue(), "application/pdf").json()

    attachment = ChatAttachment.objects.get(id=body["id"])
    assert attachment.kind == "pdf" and len(attachment.page_images) == 1
    assert attachment.extracted_text == ""


def test_digital_pdf_uses_its_text_layer(auth_client):
    from weasyprint import HTML

    pdf = HTML(string="<p>" + "The invoice total is 4,250 SAR payable within thirty days. " * 6 + "</p>").write_pdf()

    body = _upload(auth_client, "invoice.pdf", pdf, "application/pdf").json()

    attachment = ChatAttachment.objects.get(id=body["id"])
    assert "invoice total is 4,250 SAR" in attachment.extracted_text
    assert attachment.page_images == []


def test_size_limit(auth_client, monkeypatch):
    monkeypatch.setattr("chat.attachments.MAX_BYTES", 100)

    response = _upload(auth_client, "big.txt", b"x" * 500, "text/plain")

    assert response.status_code == 400


def test_attachments_reach_the_model_and_are_bound_to_the_message(auth_client, user, monkeypatch):
    image_id = _upload(auth_client, "chart.png", _png(), "image/png").json()["id"]
    note_id = _upload(auth_client, "note.txt", b"Budget: 1,200", "text/plain").json()["id"]
    script = Script(["It shows a chart."])
    use_script(monkeypatch, script)
    conversation = Conversation.objects.create(owner=user)

    response, events = stream(
        auth_client, conversation.id, content="What is this?", chat_mode="general",
        attachment_ids=[image_id, note_id],
    )

    assert response.status_code == 200
    human = script.prompts[0][-1]
    assert isinstance(human.content, list)
    text = human.content[0]["text"]
    assert text.startswith("What is this?") and 'name="note.txt"' in text and "Budget: 1,200" in text
    assert any(part.get("type") == "image_url" and part["image_url"].startswith("data:image/jpeg;base64,")
               for part in human.content)
    user_message = Message.objects.get(role="user")
    assert set(str(a.id) for a in user_message.attachments.all()) == {image_id, note_id}
    history = auth_client.get(f"/api/conversations/{conversation.id}/").json()["messages"]
    assert [a["filename"] for a in history[0]["attachments"]] == ["chart.png", "note.txt"]


def test_attachments_cannot_be_reused_or_borrowed(auth_client, other_client, user, other_user, monkeypatch):
    use_script(monkeypatch, Script(["ok", "ok"]))
    attachment_id = _upload(auth_client, "a.txt", b"hello", "text/plain").json()["id"]
    mine = Conversation.objects.create(owner=user)
    theirs = Conversation.objects.create(owner=other_user)

    borrowed, _ = stream(other_client, theirs.id, chat_mode="general", attachment_ids=[attachment_id])
    first, _ = stream(auth_client, mine.id, chat_mode="general", attachment_ids=[attachment_id])
    again, _ = stream(auth_client, mine.id, chat_mode="general", attachment_ids=[attachment_id])

    assert borrowed.status_code == 404
    assert first.status_code == 200
    assert again.status_code == 404


def test_images_never_reach_text_only_prompts():
    message = HumanMessage(content=[
        {"type": "text", "text": "Describe"},
        {"type": "image_url", "image_url": "data:image/jpeg;base64," + "A" * 5000},
    ])

    assert text_of(message) == "Describe\n[image]"


def test_text_budget_and_image_cap(tmp_path):
    long_text = ChatAttachment(filename="long.txt", kind="text", extracted_text="x" * 50_000)
    image_path = tmp_path / "i.jpg"
    Image.new("RGB", (10, 10)).save(image_path, "JPEG")
    image = ChatAttachment(filename="i.jpg", kind="image", page_images=[str(image_path)] * 3)

    content = build_human_content("q", [long_text, image], images_already_in_history=7)

    assert "truncated" in content[0]["text"]
    assert len(content[0]["text"]) < 20_000
    # One slot left in the per-conversation image budget.
    assert sum(1 for p in content if p.get("type") == "image_url") == 1
    assert any("image limit reached" in p.get("text", "") for p in content)
