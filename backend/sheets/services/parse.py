"""Read .xlsx and .csv files into header + data rows.

Real spreadsheets rarely start with a clean header on row 1: there is a title
("Grades 2024"), a blank line, a group header merged over several columns
("Grades" over Math / Physics), and merged cells down the side ("10-B" over
thirty students). This module finds the header row, flattens stacked headers,
fills merged ranges and returns JSON-safe values so the rows can be stored and
re-profiled without reading the file again.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"
SPREADSHEET_MIMES = {XLSX_MIME, CSV_MIME}

# Encodings tried in order. cp1256 is what Excel on Arabic Windows writes when
# a sheet is saved as CSV; utf-8 is strict, so it never swallows cp1256 bytes.
CSV_ENCODINGS = ("utf-8-sig", "cp1256", "latin-1")
CSV_DELIMITERS = ",;\t|"
HEADER_SEARCH_ROWS = 20
WIDTH_SAMPLE_ROWS = 50
MAX_HEADER_CHARS = 120


class SpreadsheetError(ValueError):
    """The file cannot be read as a spreadsheet, or is over a limit."""


@dataclass
class ParsedSheet:
    index: int
    name: str
    headers: list[str]
    # [(spreadsheet_row_number, [value, ...])], values aligned with headers.
    rows: list[tuple[int, list]]
    header_row: int | None = None
    header_rows: int = 1
    notes: list[str] = field(default_factory=list)


# ---- sniffing ----------------------------------------------------------------


def is_xlsx(data: bytes, max_uncompressed: int | None = None) -> bool:
    """A zip holding an Excel workbook. Raises SpreadsheetError for a zip bomb."""
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
            if "xl/workbook.xml" not in names:
                return False
            if max_uncompressed is not None:
                total = sum(info.file_size for info in archive.infolist())
                if total > max_uncompressed:
                    raise SpreadsheetError("the workbook is too large once uncompressed")
    except zipfile.BadZipFile:
        return False
    return True


def decode_text(data: bytes) -> str | None:
    if b"\x00" in data[:65536]:
        return None
    for encoding in CSV_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def sniff_delimiter(text: str) -> str | None:
    """The delimiter of a CSV with at least two columns, or None."""
    sample = "\n".join(text.splitlines()[:50])
    if not sample.strip():
        return None
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
        delimiter = dialect.delimiter
    except csv.Error:
        # The sniffer gives up on short or irregular files; count instead.
        counts = {d: sample.count(d) for d in CSV_DELIMITERS}
        delimiter = max(counts, key=counts.get)
        if counts[delimiter] == 0:
            return None
    rows = [r for r in csv.reader(io.StringIO(sample), delimiter=delimiter) if any(c.strip() for c in r)]
    if not rows or max(len(r) for r in rows) < 2:
        return None
    return delimiter


def looks_like_csv(data: bytes) -> bool:
    text = decode_text(data[:65536])
    return text is not None and sniff_delimiter(text) is not None


# ---- reading -----------------------------------------------------------------


def read_file(path: str, mime: str, *, max_rows: int, max_cols: int,
              header_overrides: dict[int, int | None] | None = None) -> list[ParsedSheet]:
    """Every non-empty sheet of the file.

    `header_overrides` maps a sheet index to a 1-based header row chosen by the
    user (None: the sheet has no header row).
    """
    with open(path, "rb") as fh:
        data = fh.read()
    if mime == XLSX_MIME:
        grids = _xlsx_grids(data)
    elif mime == CSV_MIME:
        grids = [("Sheet1", _csv_grid(data))]
    else:
        raise SpreadsheetError(f"not a spreadsheet: {mime}")

    overrides = header_overrides or {}
    sheets: list[ParsedSheet] = []
    for index, (name, grid) in enumerate(grids):
        grid = _trim(grid)
        if not grid:
            continue
        if len(grid) > max_rows + HEADER_SEARCH_ROWS:
            raise SpreadsheetError(f"sheet {name!r} has more than {max_rows} rows")
        width = max(len(r) for r in grid)
        if width > max_cols:
            raise SpreadsheetError(f"sheet {name!r} has {width} columns; the limit is {max_cols}")
        sheet = _split(index, name, grid, overrides.get(index, ...))
        if len(sheet.rows) > max_rows:
            raise SpreadsheetError(f"sheet {name!r} has {len(sheet.rows)} rows; the limit is {max_rows}")
        if sheet.rows:
            sheets.append(sheet)
    if not sheets:
        raise SpreadsheetError("the file has no data rows")
    # Keep indexes contiguous: they become document page numbers.
    for position, sheet in enumerate(sheets):
        sheet.index = position
    return sheets


def _xlsx_grids(data: bytes) -> list[tuple[str, list[list]]]:
    import openpyxl

    try:
        # Not read-only: merged ranges are only available on a full workbook.
        # data_only: the cached result of a formula, not the formula text.
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:
        raise SpreadsheetError("not a readable Excel workbook") from exc

    grids = []
    for worksheet in workbook.worksheets:
        if worksheet.sheet_state != "visible":
            continue
        grid = [list(row) for row in worksheet.iter_rows(values_only=True)]
        # A merged range keeps its value in the top-left cell only; it
        # visually applies to the whole range, so copy it there.
        for merged in worksheet.merged_cells.ranges:
            value = worksheet.cell(merged.min_row, merged.min_col).value
            for r in range(merged.min_row - 1, min(merged.max_row, len(grid))):
                row = grid[r]
                for c in range(merged.min_col - 1, merged.max_col):
                    if c < len(row):
                        row[c] = value
        grids.append((worksheet.title, [[_json_value(v) for v in row] for row in grid]))
    workbook.close()
    return grids


def _csv_grid(data: bytes) -> list[list]:
    text = decode_text(data)
    if text is None:
        raise SpreadsheetError("the CSV file is not text")
    delimiter = sniff_delimiter(text)
    if delimiter is None:
        raise SpreadsheetError("could not find the CSV columns")
    return [[_json_value(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _json_value(value):
    """A JSON scalar: str, int, float, bool or None. Dates become ISO text."""
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        return int(value) if value.is_integer() and abs(value) < 2**53 else value
    if isinstance(value, dt.datetime):
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.replace(microsecond=0).isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.replace(microsecond=0).isoformat()
    if isinstance(value, dt.timedelta):
        return str(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _trim(grid: list[list]) -> list[list]:
    """Drop trailing empty cells, and empty columns on the right."""
    rows = []
    for row in grid:
        end = len(row)
        while end and _is_empty(row[end - 1]):
            end -= 1
        rows.append(row[:end])
    while rows and not rows[-1]:
        rows.pop()
    return rows


# ---- header detection --------------------------------------------------------


def _non_empty(row: list) -> list:
    return [v for v in row if not _is_empty(v)]


def _is_label(value) -> bool:
    """Text that reads like a column name, not a number or a date."""
    if not isinstance(value, str):
        return False
    stripped = value.strip()
    return bool(stripped) and not re.fullmatch(r"[\d٠-٩.,:/\-\s%+]+", stripped)


def _looks_like_header(row: list, width: int) -> bool:
    cells = _non_empty(row)
    if len(cells) < max(2, round(width * 0.5)):
        return False
    # A title merged across the top ("Grades 2024" over every column) repeats
    # one value; a header names different things.
    if len({str(v) for v in cells}) < 2:
        return False
    labels = sum(1 for v in cells if _is_label(v))
    return labels / len(cells) >= 0.7


def detect_header_row(grid: list[list]) -> int | None:
    """0-based index of the header row in `grid`, or None."""
    width = max((len(_non_empty(r)) for r in grid[:WIDTH_SAMPLE_ROWS]), default=0)
    if width < 2:
        return None
    for index, row in enumerate(grid[:HEADER_SEARCH_ROWS]):
        if _looks_like_header(row, width) and any(_non_empty(r) for r in grid[index + 1:]):
            return index
    return None


def _has_merged_run(row: list) -> bool:
    """Adjacent equal labels: a group header merged over several columns."""
    return any(
        _is_label(a) and a == b
        for a, b in zip(row, row[1:])
    )


def _split(index: int, name: str, grid: list[list], override) -> ParsedSheet:
    width = max(len(r) for r in grid)
    grid = [r + [None] * (width - len(r)) for r in grid]

    if override is ...:
        header_at = detect_header_row(grid)
    elif override is None:
        header_at = None
    else:
        header_at = max(0, min(int(override) - 1, len(grid) - 1))

    notes: list[str] = []
    header_rows = 1
    if header_at is None:
        headers = [f"Column {i + 1}" for i in range(width)]
        first_data = 0
        notes.append("no header row found")
    else:
        top = grid[header_at]
        headers = [_header_text(v) for v in top]
        first_data = header_at + 1
        below = grid[first_data] if first_data < len(grid) else []
        if below and _has_merged_run(top) and _looks_like_header(below, width):
            # "Grades" merged over Math/Physics, with the subject names below.
            headers = [_stack(a, b) for a, b in zip(top, below)]
            first_data += 1
            header_rows = 2

    rows: list[tuple[int, list]] = []
    for offset, row in enumerate(grid[first_data:]):
        if not _non_empty(row):
            continue
        rows.append((first_data + offset + 1, row))

    # Columns with neither a header nor any data are spreadsheet furniture.
    keep = [
        c for c in range(width)
        if not headers[c].startswith("\x00") or any(not _is_empty(r[c]) for _, r in rows)
    ]
    headers = [headers[c].lstrip("\x00") for c in keep]
    rows = [(n, [r[c] for c in keep]) for n, r in rows]
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]

    return ParsedSheet(
        index=index, name=name[:255] or f"Sheet{index + 1}", headers=_unique_headers(headers), rows=rows,
        header_row=(header_at + header_rows) if header_at is not None else None,
        header_rows=header_rows, notes=notes,
    )


def _header_text(value) -> str:
    """A header cell as text; an empty one is marked so it can be dropped later."""
    if _is_empty(value):
        return "\x00"
    return str(value).strip()[:MAX_HEADER_CHARS]


def _stack(top, below) -> str:
    a, b = _header_text(top), _header_text(below)
    if a == "\x00":
        return b
    if b == "\x00" or a == b:
        return a
    return f"{a} / {b}"[:MAX_HEADER_CHARS]


def _unique_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    unique = []
    for header in headers:
        key = header.casefold()
        if key in seen:
            seen[key] += 1
            unique.append(f"{header} ({seen[key]})")
        else:
            seen[key] = 1
            unique.append(header)
    return unique
